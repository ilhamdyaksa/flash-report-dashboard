import io
import os
import json
import re
from io import BytesIO
import numpy as np
import pandas as pd
import plotly.express as px
from pptx import Presentation
from pptx.util import Inches, Pt
import streamlit as st

# ==========================================
# 1. FUNGSI UTILITAS & PERHITUNGAN
# ==========================================

def hitung_jarak_haversine_vec(lat1, lon1, lat2_series, lon2_series):
    """Perhitungan Haversine cepat berbasis vektor (NumPy)."""
    R = 6371.0
    lat1_rad, lon1_rad = np.radians(lat1), np.radians(lon1)
    lat2_rad, lon2_rad = np.radians(lat2_series.to_numpy()), np.radians(lon2_series.to_numpy())

    dlat = lat2_rad - lat1_rad
    dlon = lon2_rad - lon1_rad

    a = np.sin(dlat / 2.0) ** 2 + np.cos(lat1_rad) * np.cos(lat2_rad) * np.sin(dlon / 2.0) ** 2
    c = 2 * np.arctan2(np.sqrt(a), np.sqrt(1 - a))
    return R * c

def temukan_kolom(df, keywords):
    """Mencari nama kolom berdasarkan daftar kata kunci."""
    for col in df.columns:
        col_clean = re.sub(r'\s+', ' ', str(col).strip().lower())
        if col_clean in keywords:
            return col
    for col in df.columns:
        col_clean = str(col).strip().lower()
        if any(kw in col_clean for kw in keywords):
            return col
    return None

def clean_to_numeric(series):
    """Membersihkan format string angka, desimal, dan persen."""
    return pd.to_numeric(
        series.astype(str)
        .str.replace('%', '', regex=False)
        .str.replace(',', '.', regex=False)
        .str.strip(),
        errors='coerce'
    ).fillna(0)

def hitung_ringkasan_wilayah(df):
    """Menghitung jumlah unik Provinsi, Kabupaten, Kecamatan, dan Desa."""
    prov_col = temukan_kolom(df, ['provinsi', 'province', 'prov'])
    kab_col = temukan_kolom(df, ['kabupaten', 'regency', 'kab', 'kab/kota'])
    kec_col = temukan_kolom(df, ['kecamatan', 'district', 'kec'])
    desa_col = temukan_kolom(df, ['desa', 'kelurahan', 'village'])

    return {
        'n_prov': df[prov_col].dropna().nunique() if prov_col else 0,
        'n_kab': df[kab_col].dropna().nunique() if kab_col else 0,
        'n_kec': df[kec_col].dropna().nunique() if kec_col else 0,
        'n_desa': df[desa_col].dropna().nunique() if desa_col else 0
    }

def hitung_breakdown_site_transmisi(df):
    """Menghitung jumlah BTS USO vs 4G serta transmisi VSAT vs MW."""
    cat_col = temukan_kolom(df, ['program', 'kategori', 'category', 'jenis_bts', 'bts_type', 'tipe_bts', 'sumber_dana', 'tipe site', 'tipe_site', 'tipe'])
    trans_col = temukan_kolom(df, ['transmisi', 'transmission', 'transport', 'backhaul', 'tipe_transmisi', 'sistem_transmisi'])

    uso_sites = pd.DataFrame()
    g4_sites = pd.DataFrame()

    if cat_col and cat_col in df.columns:
        cat_str = df[cat_col].astype(str).str.upper()
        uso_sites = df[cat_str.str.contains('USO', na=False)]
        g4_sites = df[cat_str.str.contains('4G', na=False)]
    else:
        uso_mask = pd.Series(False, index=df.index)
        g4_mask = pd.Series(False, index=df.index)
        for c in df.select_dtypes(include=['object']):
            s = df[c].astype(str).str.upper()
            uso_mask |= s.str.contains('USO', na=False)
            g4_mask |= s.str.contains('4G', na=False)
        uso_sites = df[uso_mask]
        g4_sites = df[g4_mask]

    uso_vsat = uso_mw = g4_vsat = g4_mw = 0

    if trans_col and trans_col in df.columns:
        if not uso_sites.empty:
            uso_tr = uso_sites[trans_col].astype(str).str.upper()
            uso_vsat = len(uso_sites[uso_tr.str.contains('VSAT', na=False)])
            uso_mw = len(uso_sites[uso_tr.str.contains('MW|MICROWAVE', regex=True, na=False)])

        if not g4_sites.empty:
            g4_tr = g4_sites[trans_col].astype(str).str.upper()
            g4_vsat = len(g4_sites[g4_tr.str.contains('VSAT', na=False)])
            g4_mw = len(g4_sites[g4_tr.str.contains('MW|MICROWAVE', regex=True, na=False)])

    return {
        'n_uso': len(uso_sites),
        'n_4g': len(g4_sites),
        'uso_vsat': uso_vsat,
        'uso_mw': uso_mw,
        'g4_vsat': g4_vsat,
        'g4_mw': g4_mw
    }

def generate_standard_pptx_report(jenis_kejadian, detail_kejadian, waktu_kejadian, loc_label, mode_label, df_result):
    """Membaca template baku internal dan menimpa teks placeholder secara presisi."""
    template_path = "template_baku.pptx"
    
    if not os.path.exists(template_path) and os.path.exists("template_baku.pptx.pptx"):
        template_path = "template_baku.pptx.pptx"
        
    if os.path.exists(template_path):
        prs = Presentation(template_path)
    else:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        for _ in range(3):
            prs.slides.add_slide(prs.slide_layouts[6])

    status_counts = df_result['status'].astype(str).str.lower().value_counts()
    bts_down = status_counts.get('down', 0)
    bts_up = status_counts.get('up', 0)
    unmonitor = len(df_result) - bts_down - bts_up

    total_traffic_mb = df_result['traffic_mb'].sum()
    total_traffic_gb = total_traffic_mb / 1024.0
    total_period_mb = df_result['total_traffic_mb_period'].sum() if 'total_traffic_mb_period' in df_result.columns else total_traffic_mb
    total_period_gb = total_period_mb / 1024.0
    total_active_users = int(df_result['max_user'].sum())

    wilayah_info = hitung_ringkasan_wilayah(df_result)
    site_info = hitung_breakdown_site_transmisi(df_result)

    replacement_dict = {
        "{{jenis_kejadian}}": str(jenis_kejadian),
        "{{detail_kejadian}}": str(detail_kejadian),
        "{{waktu_kejadian}}": str(waktu_kejadian),
        "{{lokasi}}": str(loc_label),
        "{{mode}}": str(mode_label),
        "{{total_site}}": str(len(df_result)),
        "{{site_up}}": str(bts_up),
        "{{site_down}}": str(bts_down),
        "{{unmonitor}}": str(unmonitor),
        "{{total_traffic_gb}}": f"{total_period_gb:,.2f}",
        "{{daily_traffic_mb}}": f"{total_traffic_mb:,.2f}",
        "{{daily_traffic_gb}}": f"{total_traffic_gb:,.2f}",
        "{{active_users}}": f"{total_active_users:,}",
        "{{n_prov}}": str(wilayah_info['n_prov']),
        "{{n_kab}}": str(wilayah_info['n_kab']),
        "{{n_kec}}": str(wilayah_info['n_kec']),
        "{{n_desa}}": str(wilayah_info['n_desa']),
        "{{n_4g}}": str(site_info['n_4g']),
        "{{g4_vsat}}": str(site_info['g4_vsat']),
        "{{g4_mw}}": str(site_info['g4_mw']),
        "{{n_uso}}": str(site_info['n_uso']),
        "{{uso_vsat}}": str(site_info['uso_vsat']),
        "{{uso_mw}}": str(site_info['uso_mw'])
    }

    def ganti_teks_di_paragraf(paragraph):
        p_text = paragraph.text
        p_text = p_text.replace('\u200b', '').replace('\u200d', '')
        
        cek_replace = False
        for key in replacement_dict.keys():
            if key in p_text:
                cek_replace = True
                break
                
        if cek_replace:
            for key, val in replacement_dict.items():
                p_text = p_text.replace(key, val)
            
            if len(paragraph.runs) > 0:
                for i in range(len(paragraph.runs)):
                    paragraph.runs[i].text = "" 
                paragraph.runs[0].text = p_text 
            else:
                paragraph.text = p_text

    def proses_shape(shape):
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                ganti_teks_di_paragraf(paragraph)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        ganti_teks_di_paragraf(paragraph)
        elif shape.shape_type == 6: 
            for sub_shape in shape.shapes:
                proses_shape(sub_shape)

    for slide in prs.slides:
        for shape in slide.shapes:
            proses_shape(shape)

    buffer = BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer


# ==========================================
# 2. INISIALISASI SESSION STATE
# ==========================================
if "result_df" not in st.session_state:
    st.session_state["result_df"] = pd.DataFrame()
if "timeline_df" not in st.session_state:
    st.session_state["timeline_df"] = pd.DataFrame() # Menyimpan histori per waktu (jam/menit)
if "center_coords" not in st.session_state:
    st.session_state["center_coords"] = None
if "loc_label" not in st.session_state:
    st.session_state["loc_label"] = ""
if "mode_label" not in st.session_state:
    st.session_state["mode_label"] = ""
if "generated_pptx_bytes" not in st.session_state:
    st.session_state["generated_pptx_bytes"] = None


# ==========================================
# 3. STREAMLIT UI & LOGIC
# ==========================================
st.set_page_config(page_title="Flash Report Bencana", layout="wide")

st.title("Dashboard Laporan Dampak Kejadian")
st.caption("Kementerian Komunikasi dan Digital — BAKTI")

# --- SIDEBAR: DATA SOURCE ---
st.sidebar.header("1. Upload Data Set")

file_master = st.sidebar.file_uploader("1. Upload Master Site (Koordinat):", type=["xlsx", "xls", "csv"])
file_status = st.sidebar.file_uploader("2. Upload Laporan BAKTI OSS (Raw Data):", type=["xlsx", "xls", "csv"])

df_master = pd.DataFrame()

if file_master is not None:
    df_master = pd.read_csv(file_master) if file_master.name.endswith('.csv') else pd.read_excel(file_master)
    site_col_master = temukan_kolom(df_master, ['site id', 'site_id', 'id site', 'id_site', 'site', 'siteid'])
    
    if site_col_master:
        df_master['site_id_clean'] = df_master[site_col_master].astype(str).str.strip().str.upper()
        
        if file_status is not None:
            df_status = pd.read_csv(file_status) if file_status.name.endswith('.csv') else pd.read_excel(file_status)
            
            site_col_status = temukan_kolom(df_status, ['site id', 'site_id', 'id site', 'id_site', 'site', 'siteid'])
            avail_col = temukan_kolom(df_status, ['availability (%)', 'availability(%)', 'availability', 'avail', 'avail (%)'])
            traffic_col = temukan_kolom(df_status, ['traffic/payload (mb)', 'traffic/payload(mb)', 'traffic / payload (mb)', 'payload (mb)', 'traffic (mb)', 'payload', 'traffic', 'traffic_mb', 'payload_mb'])
            user_col = temukan_kolom(df_status, ['max active user', 'max_active_user', 'active user', 'active_user', 'user', 'max active users', 'rrc', 'max_user', 'peak user'])

            if site_col_status and avail_col and len(df_status.columns) >= 2:
                df_status['site_id_clean'] = df_status[site_col_status].astype(str).str.strip().str.upper()
                
                for col_dup in ['status', 'traffic_mb', 'total_traffic_mb_period', 'max_user']:
                    if col_dup in df_master.columns:
                        df_master = df_master.drop(columns=[col_dup])

                df_status['avail_num'] = clean_to_numeric(df_status[avail_col])
                df_status['traffic_num'] = clean_to_numeric(df_status[traffic_col]) if traffic_col else 0.0
                df_status['user_num'] = clean_to_numeric(df_status[user_col]) if user_col else 0.0

                col_a_name = df_status.columns[0]
                col_b_name = df_status.columns[1]

                combined_datetime = (
                    df_status[col_a_name].astype(str).str.strip() + " " + 
                    df_status[col_b_name].astype(str).str.strip()
                )

                df_status['parsed_time'] = pd.to_datetime(combined_datetime, errors='coerce', dayfirst=True)
                if df_status['parsed_time'].isna().all():
                    df_status['parsed_time'] = pd.to_datetime(combined_datetime, errors='coerce')

                df_status['date_str'] = df_status['parsed_time'].dt.strftime('%Y-%m-%d')
                df_status['date_str'] = df_status['date_str'].fillna(df_status[col_a_name].astype(str))

                # --- SIMPAN TIMELINE UNTUK GRAFIK TAB 2 ---
                timeline_df = df_status[['site_id_clean', 'parsed_time', 'avail_num', 'traffic_num', 'user_num']].dropna(subset=['parsed_time']).copy()
                st.session_state["timeline_df"] = timeline_df

                df_last_status = df_status.sort_values('parsed_time').groupby('site_id_clean').last().reset_index()
                df_last_status['status'] = df_last_status['avail_num'].apply(lambda x: 'Down' if x == 0 else 'Up')
                
                df_processed_status = df_last_status[['site_id_clean', 'status']].copy()

                df_status['hours_count'] = df_status.groupby(['site_id_clean', 'date_str'])[col_b_name].transform('nunique')
                df_status_valid = df_status[df_status['hours_count'] >= 24].copy()
                if df_status_valid.empty:
                    df_status_valid = df_status.copy()

                daily_per_site = df_status_valid.groupby(['site_id_clean', 'date_str']).agg(
                    total_traffic_day=('traffic_num', 'sum'),
                    peak_user_day=('user_num', 'max')
                ).reset_index()

                df_metrics = daily_per_site.groupby('site_id_clean').agg(
                    traffic_mb=('total_traffic_day', 'mean'),            
                    total_traffic_mb_period=('total_traffic_day', 'sum'), 
                    max_user=('peak_user_day', 'mean')                    
                ).reset_index()

                df_processed_status = pd.merge(df_processed_status, df_metrics, on='site_id_clean', how='left')
                cols_to_merge = ['site_id_clean', 'status', 'traffic_mb', 'total_traffic_mb_period', 'max_user']
                df_master = pd.merge(df_master, df_processed_status[cols_to_merge], on='site_id_clean', how='left')
                
                df_master['status'] = df_master['status'].fillna('Unmonitor')
                df_master['traffic_mb'] = df_master['traffic_mb'].fillna(0.0)
                df_master['total_traffic_mb_period'] = df_master['total_traffic_mb_period'].fillna(0.0)
                df_master['max_user'] = df_master['max_user'].fillna(0.0)
                
                st.sidebar.success(f"Berhasil! Data status & trend siap.")
            else:
                st.sidebar.error("Kolom yang dibutuhkan tidak ditemukan pada file RAW OSS.")
        else:
            st.sidebar.warning("File Laporan OSS belum diupload.")
            df_master['status'] = 'Unmonitor'
    else:
        st.sidebar.error("Error: Kolom 'Site ID' tidak ditemukan di file Master!")
else:
    st.sidebar.info("Silakan upload file Master Data untuk memulai.")

# --- SIDEBAR: INFO KEJADIAN & FILTER ---
st.sidebar.markdown("---")
st.sidebar.header("2. Informasi Kejadian")
jenis_kejadian = st.sidebar.text_input("Jenis Kejadian:", "Gempa Bumi")
detail_kejadian = st.sidebar.text_input("Detail (Magnitudo/Skala):", "6.2 Mag")
waktu_kejadian = st.sidebar.text_input("Waktu Kejadian:", "September 2026")

st.sidebar.markdown("---")
st.sidebar.header("3. Parameter Area Terdampak")

main_category = st.sidebar.radio(
    "Metode Filter Titik Kejadian:",
    ["Pilih Berdasarkan Area / Identitas Site", "Input Koordinat Manual"]
)

# --- PROSES PENCARIAN AREA ---
if not df_master.empty and 'status' in df_master.columns:
    lat_col = temukan_kolom(df_master, ['latitude', 'lat', 'y'])
    lon_col = temukan_kolom(df_master, ['longitude', 'long', 'lon', 'x'])

    if lat_col and lon_col:
        df_master[lat_col] = pd.to_numeric(df_master[lat_col], errors='coerce')
        df_master[lon_col] = pd.to_numeric(df_master[lon_col], errors='coerce')
        df_master = df_master.dropna(subset=[lat_col, lon_col])

        if main_category == "Pilih Berdasarkan Area / Identitas Site":
            kategori_area = st.sidebar.selectbox("Pilih Kategori Pencarian:", ["Kabupaten", "Kecamatan", "Desa", "Provinsi", "Site ID"])
            key_mapping = {
                "Provinsi": temukan_kolom(df_master, ['provinsi', 'province', 'prov']),
                "Kabupaten": temukan_kolom(df_master, ['kabupaten', 'regency', 'kab']),
                "Kecamatan": temukan_kolom(df_master, ['kecamatan', 'district', 'kec']),
                "Desa": temukan_kolom(df_master, ['desa', 'kelurahan', 'village']),
                "Site ID": 'site_id_clean'
            }
            target_col = key_mapping.get(kategori_area)

            if target_col and target_col in df_master.columns:
                list_options = sorted(df_master[target_col].dropna().astype(str).unique().tolist())
                selected_val = st.sidebar.selectbox(f"Cari/Pilih {kategori_area}:", options=["-- Pilih --"] + list_options)

                filter_mode = st.sidebar.radio("Metode Cakupan Area:", ["Batas Administrasi Murni (Eksak)", "Radius Jarak (KM) dari Pusat Area"])
                radius_km = 10
                if filter_mode == "Radius Jarak (KM) dari Pusat Area":
                    radius_km = st.sidebar.slider("Radius Terdampak (KM):", 1, 100, 10)

                if st.sidebar.button("Cari & Hitung Dampak") and selected_val != "-- Pilih --":
                    matched_df = df_master[df_master[target_col].astype(str).str.lower() == selected_val.lower()]
                    if not matched_df.empty:
                        center_lat = matched_df[lat_col].mean()
                        center_lon = matched_df[lon_col].mean()

                        if filter_mode == "Batas Administrasi Murni (Eksak)":
                            final_result = matched_df.copy()
                            mode_desc = f"Batas Administrasi Murni ({kategori_area}: {selected_val})"
                        else:
                            df_master['jarak_km'] = hitung_jarak_haversine_vec(center_lat, center_lon, df_master[lat_col], df_master[lon_col])
                            final_result = df_master[df_master['jarak_km'] <= radius_km].copy()
                            mode_desc = f"Radius {radius_km} KM dari Pusat {selected_val}"

                        st.session_state["center_coords"] = (center_lat, center_lon)
                        st.session_state["result_df"] = final_result
                        st.session_state["loc_label"] = f"{kategori_area} {selected_val}"
                        st.session_state["mode_label"] = mode_desc
                        st.session_state["generated_pptx_bytes"] = None
                    else:
                        st.error("Data tidak ditemukan.")

        elif main_category == "Input Koordinat Manual":
            col_l, col_r = st.sidebar.columns(2)
            with col_l: input_lat = st.number_input("Latitude:", format="%.6f", value=-4.512300)
            with col_r: input_lon = st.number_input("Longitude:", format="%.6f", value=140.412300)
            radius_km = st.sidebar.slider("Radius Terdampak (KM):", 1, 100, 10)

            if st.sidebar.button("Hitung Area Terdampak"):
                df_master['jarak_km'] = hitung_jarak_haversine_vec(input_lat, input_lon, df_master[lat_col], df_master[lon_col])
                st.session_state["center_coords"] = (input_lat, input_lon)
                st.session_state["result_df"] = df_master[df_master['jarak_km'] <= radius_km].copy()
                st.session_state["loc_label"] = f"Koordinat ({input_lat}, {input_lon})"
                st.session_state["mode_label"] = f"Radius {radius_km} KM"
                st.session_state["generated_pptx_bytes"] = None


# --- TAMPILAN DASHBOARD UTAMA ---
result_df = st.session_state["result_df"]
center_coords = st.session_state["center_coords"]
loc_label = st.session_state["loc_label"]
mode_label = st.session_state["mode_label"]

if not result_df.empty:
    st.success(f"**Ringkasan Laporan:** {jenis_kejadian} - {detail_kejadian} | {waktu_kejadian}")
    st.markdown(f"📍 **Lokasi Fokus:** {loc_label} | **Metode:** {mode_label}")
    
    wilayah_info = hitung_ringkasan_wilayah(result_df)
    site_info = hitung_breakdown_site_transmisi(result_df)

    status_counts = result_df['status'].astype(str).str.lower().value_counts()
    bts_down = status_counts.get('down', 0)
    bts_up = status_counts.get('up', 0)
    unmonitor = len(result_df) - bts_down - bts_up 

    total_traffic_mb = result_df['traffic_mb'].sum()
    total_traffic_gb = total_traffic_mb / 1024.0
    total_period_mb = result_df['total_traffic_mb_period'].sum() if 'total_traffic_mb_period' in result_df.columns else total_traffic_mb
    total_period_gb = total_period_mb / 1024.0
    total_active_users = int(result_df['max_user'].sum())

    st.subheader("⚡ Status Jaringan Terdampak (Jam Terakhir)")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Total Site Terdampak", len(result_df))
    c2.metric("Site Up", bts_up)
    c3.metric("Site Down", bts_down, delta="- Kritis" if bts_down > 0 else "Normal", delta_color="inverse")
    c4.metric("Unmonitor", unmonitor)

    st.markdown("---")
    st.subheader("📡 Kategori & Transmisi Site BAKTI")
    col_4g, col_uso = st.columns(2)
    with col_4g:
        st.markdown("#### **Program BTS 4G**")
        m_4g_total, m_4g_vsat, m_4g_mw = st.columns(3)
        m_4g_total.metric("Total Site", site_info['n_4g'])
        m_4g_vsat.metric("Transmisi VSAT", site_info['g4_vsat'])
        m_4g_mw.metric("Transmisi MW", site_info['g4_mw'])
    with col_uso:
        st.markdown("#### **Program BTS USO**")
        m_uso_total, m_uso_vsat, m_uso_mw = st.columns(3)
        m_uso_total.metric("Total Site", site_info['n_uso'])
        m_uso_vsat.metric("Transmisi VSAT", site_info['uso_vsat'])
        m_uso_mw.metric("Transmisi MW", site_info['uso_mw'])

    st.divider()
    
    # 3 TAB (Peta, Trend, Tabel)
    tab1, tab2, tab3 = st.tabs(["🗺️ Visual Peta Sebaran", "📈 Tren Layanan & Traffic", "📋 Detail Data Site"])
    
    with tab1:
        st.markdown("#### **Peta Sebaran Operasional (Auto-Zoom)**")
        plotly_df = result_df.copy()
        plotly_df['Status_Site'] = plotly_df['status'].astype(str).str.capitalize()
        color_map = {'Up': '#00BFFF', 'Down': '#FF5252', 'Unmonitor': '#FFC107'}

        mean_lat = float(plotly_df[lat_col].mean())
        mean_lon = float(plotly_df[lon_col].mean())

        fig_map = px.scatter_map(
            plotly_df, lat=lat_col, lon=lon_col, color='Status_Site',
            color_discrete_map=color_map, hover_name='site_id_clean',
            hover_data={'status': True, 'traffic_mb': ':.2f', 'max_user': True, lat_col: False, lon_col: False},
            zoom=8, center=dict(lat=mean_lat, lon=mean_lon), height=550
        )

        if center_coords:
            center_marker_df = pd.DataFrame({'lat': [center_coords[0]], 'lon': [center_coords[1]]})
            fig_map.add_scattermap(
                lat=center_marker_df['lat'], lon=center_marker_df['lon'],
                mode='markers', marker=dict(size=18, color='darkred', symbol='star'),
                name=f"Pusat {jenis_kejadian}"
            )

        map_layers_list = [{"below": 'traces', "sourcetype": "raster", "source": ["https://server.arcgisonline.com/ArcGIS/rest/services/World_Imagery/MapServer/tile/{z}/{y}/{x}"]}]

        geojson_path = 'indonesia.geojson'
        if os.path.exists(geojson_path):
            file_size_mb = os.path.getsize(geojson_path) / (1024 * 1024)
            if file_size_mb <= 15.0:
                try:
                    with open(geojson_path, 'r', encoding='utf-8') as f:
                        geojson_kab = json.load(f)
                    map_layers_list.append({"sourcetype": "geojson", "source": geojson_kab, "type": "line", "color": "rgba(255, 255, 255, 0.6)", "line": {"width": 1.5}, "below": "traces"})
                except Exception:
                    pass

        fig_map.update_layout(map_style="white-bg", map_layers=map_layers_list, margin={"r":0,"t":0,"l":0,"b":0})
        st.plotly_chart(fig_map, use_container_width=True)
        
    with tab2:
        st.markdown("#### **Tren Layanan Historis pada Area Terdampak**")
        timeline_df = st.session_state.get("timeline_df", pd.DataFrame())
        
        if not timeline_df.empty:
            affected_sites = result_df['site_id_clean'].unique()
            trend_df = timeline_df[timeline_df['site_id_clean'].isin(affected_sites)]
            
            if not trend_df.empty:
                trend_grouped = trend_df.groupby('parsed_time').agg(
                    avg_avail=('avail_num', 'mean'),
                    total_traffic=('traffic_num', 'sum'),
                    total_user=('user_num', 'sum')
                ).reset_index()
                
                trend_grouped = trend_grouped.sort_values('parsed_time')
                trend_grouped['total_traffic_gb'] = trend_grouped['total_traffic'] / 1024.0
                
                col_c1, col_c2 = st.columns(2)
                
                fig_avail = px.line(
                    trend_grouped, x='parsed_time', y='avg_avail', 
                    title='📈 Tren Rata-rata Availability (%)',
                    labels={'parsed_time': 'Waktu', 'avg_avail': 'Availability (%)'},
                    markers=True
                )
                fig_avail.update_traces(line_color='#2E7D32', marker=dict(size=4))
                fig_avail.update_layout(xaxis_tickangle=-45, plot_bgcolor='white', yaxis=dict(gridcolor='lightgray'))
                
                fig_traffic = px.line(
                    trend_grouped, x='parsed_time', y='total_traffic_gb', 
                    title='📊 Tren Total Traffic (GB)',
                    labels={'parsed_time': 'Waktu', 'total_traffic_gb': 'Traffic (GB)'},
                    markers=True
                )
                fig_traffic.update_traces(line_color='#0277BD', marker=dict(size=4))
                fig_traffic.update_layout(xaxis_tickangle=-45, plot_bgcolor='white', yaxis=dict(gridcolor='lightgray'))
                
                fig_user = px.line(
                    trend_grouped, x='parsed_time', y='total_user', 
                    title='👥 Tren Total Active Users',
                    labels={'parsed_time': 'Waktu', 'total_user': 'Total User'},
                    markers=True
                )
                fig_user.update_traces(line_color='#F9A825', marker=dict(size=4))
                fig_user.update_layout(xaxis_tickangle=-45, plot_bgcolor='white', yaxis=dict(gridcolor='lightgray'))
                
                with col_c1:
                    st.plotly_chart(fig_avail, use_container_width=True)
                    st.plotly_chart(fig_user, use_container_width=True)
                with col_c2:
                    st.plotly_chart(fig_traffic, use_container_width=True)
                    
            else:
                st.info("Tidak ada riwayat pergerakan data pada site di area ini.")
        else:
            st.info("Riwayat timeline tidak tersedia. Pastikan format waktu pada raw data terbaca.")

    with tab3:
        st.dataframe(result_df.drop(columns=['jarak_km', 'site_id_clean'], errors='ignore'), use_container_width=True)

    st.markdown("---")
    st.subheader("📄 Laporan PowerPoint (PPTX)")

    col_gen, col_down = st.columns([1, 1])

    with col_gen:
        if st.button("🚀 Generate PPT Flash Report", type="primary", use_container_width=True):
            with st.spinner("Memproses pembuatan slide PowerPoint..."):
                st.session_state["generated_pptx_bytes"] = generate_standard_pptx_report(
                    jenis_kejadian, detail_kejadian, waktu_kejadian,
                    loc_label, mode_label, result_df
                )
            st.success("File PPTX Berhasil Dibuat!")

    with col_down:
        if st.session_state["generated_pptx_bytes"] is not None:
            # Mengamankan nama lokasi dari karakter '/' atau ':' agar tidak error saat disimpan di Windows/Mac
            safe_lokasi = loc_label.replace('/', '_').replace(':', '_')
            
            st.download_button(
                label="📥 Download File PPTX",
                data=st.session_state["generated_pptx_bytes"],
                file_name=f"Flash Report_{jenis_kejadian}_{safe_lokasi}.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                use_container_width=True
            )
        else:
            st.info("Klik tombol **🚀 Generate PPT Flash Report** di sebelah kiri terlebih dahulu.")